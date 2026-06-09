"""Export WBS and RAID artifacts using APS1049 Deloitte deliverable templates."""

from __future__ import annotations

import io
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.worksheet.worksheet import Worksheet

from workplan_agent.display_values import Notes, display_cell
from workplan_agent.files.pdf_render import (
    add_landscape_table_section,
    pdf_text,
    render_table,
    write_body_text,
    write_brief_header,
    write_bullet_list,
    write_section_heading,
)
from workplan_agent.models.brief import ParsedBrief
from workplan_agent.models.raid import ActionEntry, DecisionEntry, IssueEntry, RAIDLog, RiskEntry
from workplan_agent.models.raci import RaciLetter, RaciMatrix
from workplan_agent.models.status_report import WeeklyStatusReport
from workplan_agent.models.wbs import WBS, WorkPackage

_PACKAGE_TEMPLATES = Path(__file__).resolve().parent.parent / "templates" / "deloitte"
RAID_TEMPLATE_PATH = _PACKAGE_TEMPLATES / "raid_log_template.xlsx"
RACI_TEMPLATE_PATH = _PACKAGE_TEMPLATES / "raci_template.xlsx"

_RACI_SHEET = "RACI Table Sample"
_RACI_HEADER_ROW = 4
_RACI_DATA_START_ROW = 5

_DETAIL_SHEET = "RAID Log - Detail"
_DETAIL_DATA_START_ROW = 7


def _confirmed_marker_suffix(item: object) -> str:
    return "*" if getattr(item, "from_confirmed_answer", False) else ""


def resolve_raid_template(path: Path | None = None) -> Path:
    if path is not None:
        if not path.is_file():
            raise FileNotFoundError(f"RAID template not found: {path}")
        return path
    if RAID_TEMPLATE_PATH.is_file():
        return RAID_TEMPLATE_PATH
    repo_root = Path.cwd()
    for candidate in (
        repo_root / "APS1049_Team 2-RAID Log-Template .xlsx",
        repo_root / "templates" / "deloitte" / "raid_log_template.xlsx",
    ):
        if candidate.is_file():
            return candidate
    raise FileNotFoundError(
        "RAID Excel template not found. Expected bundled template or "
        "APS1049_Team 2-RAID Log-Template .xlsx in the project root."
    )


def export_wbs_deloitte_docx(wbs: WBS, parsed_brief: ParsedBrief) -> bytes:
    """Build a WBS Word document matching the Deloitte table layout."""
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError as exc:
        raise RuntimeError("DOCX export is not installed.") from exc

    document = Document()
    document.add_heading("Work Breakdown Structure", level=0)
    document.add_paragraph(f"Project: {parsed_brief.project_title}")
    document.add_paragraph(f"Client: {parsed_brief.client_name}")

    headers = [
        "Level",
        "WBS Code",
        "WBS Element Title",
        "Description / Scope of Work",
        "Key Deliverable",
        "Owner",
        "Duration",
        "Dependencies",
    ]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for col_idx, header in enumerate(headers):
        table.rows[0].cells[col_idx].text = header

    for wp in sorted(wbs.work_packages, key=_wbs_sort_key):
        row = table.add_row().cells
        row[0].text = str(wp.level)
        row[1].text = f"{wp.wbs_code}{_confirmed_marker_suffix(wp)}"
        row[2].text = display_cell(wp.wbs_element_title)
        row[3].text = display_cell(wp.description_scope, note=Notes.SIGNOFF_DETAIL)
        row[4].text = display_cell(wp.key_deliverable, note=Notes.SIGNOFF)
        row[5].text = display_cell(wp.owner, note=Notes.SIGNOFF_ASSIGNED)
        row[6].text = display_cell(wp.estimated_duration, note=Notes.SIGNOFF_ESTIMATE)
        row[7].text = display_cell(wp.dependencies, note=Notes.NONE_IDENTIFIED)

    for para in document.paragraphs:
        for run in para.runs:
            run.font.size = Pt(11)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


_WBS_PDF_HEADERS = [
    "Level",
    "WBS Code",
    "WBS Element Title",
    "Description / Scope of Work",
    "Key Deliverable",
    "Owner",
    "Duration",
    "Dependencies",
]

# Landscape A4 effective width with 10 mm side margins (mm).
_WBS_PDF_COL_WIDTHS = (14, 20, 34, 58, 40, 32, 26, 44)

_RISK_PDF_HEADERS = [
    "ID",
    "Title",
    "Description",
    "Status",
    "Mitigation",
    "Owner",
    "Owner Side",
    "WBS Links",
    "Probability",
    "Impact",
    "Risk Rating",
    "Risk Appetite",
    "Priority",
]
_RISK_PDF_COL_WIDTHS = (12, 24, 38, 16, 30, 22, 18, 20, 18, 16, 18, 20, 20)

_ACTION_PDF_HEADERS = [
    "ID",
    "Title",
    "Description",
    "Status",
    "Owner",
    "Owner Side",
    "WBS Links",
    "Priority",
]
_ACTION_PDF_COL_WIDTHS = (12, 28, 55, 18, 25, 20, 22, 18)

_ISSUE_PDF_HEADERS = [
    "ID",
    "Title",
    "Description",
    "Resolution",
    "Owner",
    "Owner Side",
    "Priority",
]
_ISSUE_PDF_COL_WIDTHS = (12, 30, 55, 35, 25, 20, 18)

_DECISION_PDF_HEADERS = [
    "ID",
    "Title",
    "Description / Justification",
    "Resolution",
    "Owner",
    "Owner Side",
    "Priority",
]
_DECISION_PDF_COL_WIDTHS = (12, 30, 55, 35, 25, 20, 18)


_MILESTONE_PDF_HEADERS = [
    "Milestone",
    "Target Date",
    "Actual Date",
    "Owner Side",
    "Status",
]
_MILESTONE_PDF_COL_WIDTHS = (40, 35, 35, 28, 30)


def export_wbs_deloitte_pdf(wbs: WBS, parsed_brief: ParsedBrief) -> bytes:
    """Build a landscape WBS PDF with a proper multi-column table."""
    try:
        from fpdf import FPDF
    except ImportError as exc:
        raise RuntimeError("PDF export is not installed.") from exc

    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_margins(10, 12, 10)
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()
    write_brief_header(pdf, "Work Breakdown Structure", parsed_brief)

    phases = [wp for wp in sorted(wbs.work_packages, key=_wbs_sort_key) if wp.level == 1]
    if phases:
        write_section_heading(pdf, "Phases (level 1)")
        pdf.set_font("Helvetica", size=10)
        for wp in phases:
            pdf.cell(
                0,
                5,
                pdf_text(f"  - {wp.wbs_code} - {wp.wbs_element_title}"),
                new_x="LMARGIN",
                new_y="NEXT",
            )
        pdf.ln(2)

    write_section_heading(pdf, "Work packages")
    render_table(
        pdf,
        _WBS_PDF_HEADERS,
        [
            [
                display_cell(wp.level),
                display_cell(f"{wp.wbs_code}{_confirmed_marker_suffix(wp)}"),
                display_cell(wp.wbs_element_title),
                display_cell(wp.description_scope, note=Notes.SIGNOFF_DETAIL),
                display_cell(wp.key_deliverable, note=Notes.SIGNOFF),
                display_cell(wp.owner, note=Notes.SIGNOFF_ASSIGNED),
                display_cell(wp.estimated_duration, note=Notes.SIGNOFF_ESTIMATE),
                display_cell(wp.dependencies, note=Notes.NONE_IDENTIFIED),
            ]
            for wp in sorted(wbs.work_packages, key=_wbs_sort_key)
        ],
        col_widths=_WBS_PDF_COL_WIDTHS,
    )

    return bytes(pdf.output())


def export_raid_deloitte_pdf(raid: RAIDLog, parsed_brief: ParsedBrief) -> bytes:
    """Build a RAID PDF with landscape tables per section."""
    try:
        from fpdf import FPDF
    except ImportError as exc:
        raise RuntimeError("PDF export is not installed.") from exc

    pdf = FPDF(unit="mm", format="A4")
    pdf.set_margins(18, 18, 18)
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()
    write_brief_header(pdf, "RAID Log", parsed_brief)

    risk_rows = [
        [
            display_cell(f"{item.id}{_confirmed_marker_suffix(item)}"),
            display_cell(item.title),
            display_cell(item.description, note=Notes.SIGNOFF_DETAIL),
            display_cell(item.status),
            display_cell(item.mitigation, note=Notes.NOT_YET_DEFINED),
            display_cell(item.owner, note=Notes.SIGNOFF_ASSIGNED),
            display_cell(item.owner_side),
            display_cell(item.wbs_links, note=Notes.NOT_LINKED_WBS),
            display_cell(item.probability),
            display_cell(item.impact),
            display_cell(item.risk_rating),
            display_cell(item.risk_appetite),
            display_cell(item.priority),
        ]
        for item in raid.risks
    ]
    if risk_rows:
        add_landscape_table_section(
            pdf,
            "Risks",
            _RISK_PDF_HEADERS,
            risk_rows,
            col_widths=_RISK_PDF_COL_WIDTHS,
        )

    action_rows = [
        [
            display_cell(f"{item.id}{_confirmed_marker_suffix(item)}"),
            display_cell(item.title),
            display_cell(item.description, note=Notes.SIGNOFF_DETAIL),
            display_cell(item.status),
            display_cell(item.owner, note=Notes.SIGNOFF_ASSIGNED),
            display_cell(item.owner_side),
            display_cell(item.wbs_links, note=Notes.NOT_LINKED_WBS),
            display_cell(item.priority),
        ]
        for item in raid.actions
    ]
    if action_rows:
        add_landscape_table_section(
            pdf,
            "Actions",
            _ACTION_PDF_HEADERS,
            action_rows,
            col_widths=_ACTION_PDF_COL_WIDTHS,
        )

    if raid.issues:
        issue_rows = [
            [
                display_cell(f"{item.id}{_confirmed_marker_suffix(item)}"),
                display_cell(item.title),
                display_cell(item.description, note=Notes.SIGNOFF_DETAIL),
                display_cell(item.resolution, note=Notes.PENDING_RESOLUTION),
                display_cell(item.owner, note=Notes.SIGNOFF_ASSIGNED),
                display_cell(item.owner_side),
                display_cell(item.priority),
            ]
            for item in raid.issues
        ]
    else:
        issue_rows = [
            [
                "—",
                "No issues identified",
                display_cell(None, note=Notes.NONE_IDENTIFIED),
                display_cell(None, note=Notes.PENDING_RESOLUTION),
                display_cell(None, note=Notes.SIGNOFF_ASSIGNED),
                display_cell(None),
                display_cell(None),
            ]
        ]
    add_landscape_table_section(
        pdf,
        "Issues",
        _ISSUE_PDF_HEADERS,
        issue_rows,
        col_widths=_ISSUE_PDF_COL_WIDTHS,
    )

    decision_rows = [
        [
            display_cell(f"{item.id}{_confirmed_marker_suffix(item)}"),
            display_cell(item.title),
            display_cell(item.description_justification, note=Notes.SIGNOFF_DETAIL),
            display_cell(item.resolution, note=Notes.PENDING_RESOLUTION),
            display_cell(item.owner, note=Notes.SIGNOFF_ASSIGNED),
            display_cell(item.owner_side),
            display_cell(item.priority),
        ]
        for item in raid.decisions
    ]
    if decision_rows:
        add_landscape_table_section(
            pdf,
            "Decisions",
            _DECISION_PDF_HEADERS,
            decision_rows,
            col_widths=_DECISION_PDF_COL_WIDTHS,
        )

    return bytes(pdf.output())


def resolve_raci_template(path: Path | None = None) -> Path:
    if path is not None:
        if not path.is_file():
            raise FileNotFoundError(f"RACI template not found: {path}")
        return path
    if RACI_TEMPLATE_PATH.is_file():
        return RACI_TEMPLATE_PATH
    repo_root = Path.cwd()
    candidate = repo_root / "templates" / "deloitte" / "raci_template.xlsx"
    if candidate.is_file():
        return candidate
    raise FileNotFoundError(
        "RACI Excel template not found. Expected bundled raci_template.xlsx."
    )


def export_raci_deloitte_xlsx(
    raci_matrix: RaciMatrix,
    parsed_brief: ParsedBrief,
    template_path: Path | None = None,
) -> bytes:
    """Fill the RACI Excel template."""
    path = resolve_raci_template(template_path)
    workbook = load_workbook(path)
    if _RACI_SHEET not in workbook.sheetnames:
        raise ValueError(f"RACI template missing required sheet: {_RACI_SHEET!r}")
    ws = workbook[_RACI_SHEET]
    project = parsed_brief.project_title or "Project"
    _set_cell(ws, 2, 2, f"RACI — {project}")
    _set_cell(ws, _RACI_HEADER_ROW, 2, "WBS Code")
    _set_cell(ws, _RACI_HEADER_ROW, 3, "Task / Activity")
    for col_idx, role in enumerate(raci_matrix.roles, start=4):
        _set_cell(ws, _RACI_HEADER_ROW, col_idx, role)

    end_row = max(_RACI_DATA_START_ROW + len(raci_matrix.rows) + 5, ws.max_row)
    _clear_row_range(ws, _RACI_DATA_START_ROW, end_row, max(3 + len(raci_matrix.roles), 8))

    for row_idx, row in enumerate(raci_matrix.rows, start=_RACI_DATA_START_ROW):
        _set_cell(ws, row_idx, 2, row.wbs_code)
        _set_cell(
            ws,
            row_idx,
            3,
            f"{row.activity}{_confirmed_marker_suffix(row)}",
        )
        for col_idx, role in enumerate(raci_matrix.roles, start=4):
            letter = row.assignments.get(role)
            value = letter.value if isinstance(letter, RaciLetter) else letter
            _set_cell(ws, row_idx, col_idx, value)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_raci_deloitte_pdf(
    raci_matrix: RaciMatrix,
    parsed_brief: ParsedBrief,
) -> bytes:
    """Build a landscape RACI matrix PDF."""
    try:
        from fpdf import FPDF
    except ImportError as exc:
        raise RuntimeError("PDF export is not installed.") from exc

    pdf = FPDF(unit="mm", format="A4")
    pdf.set_margins(18, 18, 18)
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()
    write_brief_header(pdf, "RACI Matrix", parsed_brief)

    headers = ["WBS Code", "Task / Activity", *raci_matrix.roles]
    col_count = len(headers)
    col_width = max(180 / col_count, 18)
    col_widths = tuple([col_width, max(col_width * 1.5, 35), *[col_width] * (col_count - 2)])

    rows = [
        [
            display_cell(row.wbs_code),
            display_cell(f"{row.activity}{_confirmed_marker_suffix(row)}"),
        ]
        + [
            display_cell(
                row.assignments[role].value
                if role in row.assignments
                else None
            )
            for role in raci_matrix.roles
        ]
        for row in raci_matrix.rows
    ]
    pdf.add_page(orientation="L")
    pdf.set_margins(10, 12, 10)
    render_table(pdf, headers, rows, col_widths=col_widths)
    return bytes(pdf.output())


def export_status_deloitte_pdf(
    status_report: WeeklyStatusReport,
    parsed_brief: ParsedBrief,
) -> bytes:
    """Build a weekly status report PDF."""
    try:
        from fpdf import FPDF
    except ImportError as exc:
        raise RuntimeError("PDF export is not installed.") from exc

    pdf = FPDF(unit="mm", format="A4")
    pdf.set_margins(18, 18, 18)
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()
    width = pdf.epw

    pdf.set_font("Helvetica", style="B", size=16)
    pdf.cell(0, 9, pdf_text("Weekly Status Report"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=11)
    for line in (
        f"Project: {status_report.project_name}",
        f"Client: {parsed_brief.client_name}",
        f"Date: {status_report.date}",
        f"Project Manager: {status_report.project_manager}",
    ):
        pdf.cell(0, 6, pdf_text(line), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)

    write_section_heading(pdf, "Summary")
    write_body_text(pdf, status_report.summary, width=width)

    write_section_heading(pdf, "New Initiatives")
    if status_report.new_initiatives:
        write_bullet_list(pdf, status_report.new_initiatives, width=width)
    else:
        write_body_text(pdf, "No new initiatives.", width=width)

    write_section_heading(pdf, "Blockers")
    if status_report.blockers:
        write_bullet_list(pdf, status_report.blockers, width=width)
    else:
        write_body_text(pdf, "No blockers.", width=width)

    write_section_heading(pdf, "Risks")
    if status_report.risks:
        write_bullet_list(pdf, status_report.risks, width=width)
    else:
        write_body_text(pdf, "No risks.", width=width)

    write_section_heading(pdf, "Milestones")
    milestone_rows = [
        [
            display_cell(f"{milestone.name}{_confirmed_marker_suffix(milestone)}"),
            display_cell(milestone.target_date),
            display_cell(milestone.actual_date, note=Notes.NOT_YET_REACHED),
            display_cell(milestone.owner_side),
            display_cell(milestone.status),
        ]
        for milestone in status_report.milestones
    ]
    if milestone_rows:
        render_table(
            pdf,
            _MILESTONE_PDF_HEADERS,
            milestone_rows,
            col_widths=_MILESTONE_PDF_COL_WIDTHS,
            font_size=9,
            line_height=5,
        )
    else:
        write_body_text(pdf, "No milestones.", width=width)

    write_section_heading(pdf, "Supporting Links")
    if status_report.supporting_links:
        write_bullet_list(pdf, status_report.supporting_links, width=width)
    else:
        write_body_text(pdf, "No supporting links.", width=width)

    return bytes(pdf.output())


def export_raid_deloitte_xlsx(
    raid: RAIDLog,
    parsed_brief: ParsedBrief,
    template_path: Path | None = None,
) -> bytes:
    """Fill the RAID Excel template detail sheet only."""
    path = resolve_raid_template(template_path)
    workbook = load_workbook(path)
    project = parsed_brief.project_title or "Project"
    today = date.today().strftime("%Y-%m-%d")

    if _DETAIL_SHEET not in workbook.sheetnames:
        raise ValueError(f"RAID template missing required sheet: {_DETAIL_SHEET!r}")
    _fill_raid_detail_sheet(workbook[_DETAIL_SHEET], raid, project, today)
    for sheet_name in list(workbook.sheetnames):
        if sheet_name != _DETAIL_SHEET:
            workbook.remove(workbook[sheet_name])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _wbs_sort_key(wp: WorkPackage) -> tuple[int, ...]:
    parts: list[int] = []
    for segment in wp.wbs_code.replace("-", ".").split("."):
        if segment.isdigit():
            parts.append(int(segment))
        else:
            parts.append(0)
    return tuple(parts)


def _writable_cell(ws: Worksheet, row: int, col: int):
    """Return the cell that accepts writes (top-left of a merged range)."""
    cell = ws.cell(row=row, column=col)
    if not isinstance(cell, MergedCell):
        return cell
    for merged in ws.merged_cells.ranges:
        if (
            merged.min_row <= row <= merged.max_row
            and merged.min_col <= col <= merged.max_col
        ):
            return ws.cell(row=merged.min_row, column=merged.min_col)
    return cell


def _set_cell(
    ws: Worksheet,
    row: int,
    col: int,
    value: object,
    *,
    note: str | None = None,
) -> None:
    _writable_cell(ws, row, col).value = display_cell(value, note=note)


def _clear_row_range(ws: Worksheet, start_row: int, end_row: int, max_col: int) -> None:
    cleared_merges: set[str] = set()
    for row in range(start_row, end_row + 1):
        for col in range(1, max_col + 1):
            cell = ws.cell(row=row, column=col)
            if isinstance(cell, MergedCell):
                for merged in ws.merged_cells.ranges:
                    if (
                        merged.min_row <= row <= merged.max_row
                        and merged.min_col <= col <= merged.max_col
                    ):
                        merge_key = str(merged)
                        if merge_key not in cleared_merges:
                            cleared_merges.add(merge_key)
                            ws.cell(
                                row=merged.min_row,
                                column=merged.min_col,
                                value=None,
                            )
                        break
            else:
                cell.value = None


def _fill_raid_detail_sheet(
    ws: Worksheet,
    raid: RAIDLog,
    project: str,
    today: str,
) -> None:
    _set_cell(ws, 4, 3, project)
    _set_cell(ws, 1, 9, today)

    max_items = max(
        len(raid.risks),
        len(raid.actions),
        len(raid.issues),
        len(raid.decisions),
        1,
    )
    end_row = _DETAIL_DATA_START_ROW + max(max_items, 6) + 5
    _clear_row_range(ws, _DETAIL_DATA_START_ROW, end_row, 34)

    for index in range(max_items):
        row = _DETAIL_DATA_START_ROW + index
        if index < len(raid.risks):
            _write_detail_risk(ws, row, raid.risks[index])
        if index < len(raid.actions):
            _write_detail_action(ws, row, raid.actions[index])
        if index < len(raid.issues):
            _write_detail_issue(ws, row, raid.issues[index])
        if index < len(raid.decisions):
            _write_detail_decision(ws, row, raid.decisions[index])


def _write_detail_risk(ws: Worksheet, row: int, item: RiskEntry) -> None:
    _set_cell(ws, row, 1, f"{item.id}{_confirmed_marker_suffix(item)}")
    _set_cell(ws, row, 2, item.title)
    _set_cell(ws, row, 3, item.description, note=Notes.SIGNOFF_DETAIL)
    _set_cell(ws, row, 4, item.status)
    _set_cell(ws, row, 5, item.mitigation, note=Notes.NOT_YET_DEFINED)
    _set_cell(ws, row, 6, item.owner, note=Notes.SIGNOFF_ASSIGNED)
    _set_cell(ws, row, 7, item.wbs_links, note=Notes.NOT_LINKED_WBS)
    _set_cell(ws, row, 8, item.probability)
    _set_cell(ws, row, 9, item.impact)
    _set_cell(ws, row, 10, item.risk_rating)
    _set_cell(ws, row, 11, item.risk_appetite)
    _set_cell(ws, row, 12, item.priority)


def _write_detail_action(ws: Worksheet, row: int, item: ActionEntry) -> None:
    _set_cell(ws, row, 14, f"{item.id}{_confirmed_marker_suffix(item)}")
    _set_cell(ws, row, 15, item.title)
    _set_cell(ws, row, 16, item.description, note=Notes.SIGNOFF_DETAIL)
    _set_cell(ws, row, 17, item.status)
    _set_cell(ws, row, 18, item.owner, note=Notes.SIGNOFF_ASSIGNED)
    _set_cell(ws, row, 19, item.wbs_links, note=Notes.NOT_LINKED_WBS)
    _set_cell(ws, row, 20, item.priority)


def _write_detail_issue(ws: Worksheet, row: int, item: IssueEntry) -> None:
    _set_cell(ws, row, 22, f"{item.id}{_confirmed_marker_suffix(item)}")
    _set_cell(ws, row, 23, item.title)
    _set_cell(ws, row, 24, item.description, note=Notes.SIGNOFF_DETAIL)
    _set_cell(ws, row, 25, item.resolution, note=Notes.PENDING_RESOLUTION)
    _set_cell(ws, row, 26, item.owner, note=Notes.SIGNOFF_ASSIGNED)
    _set_cell(ws, row, 27, item.priority)


def _write_detail_decision(ws: Worksheet, row: int, item: DecisionEntry) -> None:
    _set_cell(ws, row, 29, f"{item.id}{_confirmed_marker_suffix(item)}")
    _set_cell(ws, row, 30, item.title)
    _set_cell(ws, row, 31, item.description_justification, note=Notes.SIGNOFF_DETAIL)
    _set_cell(ws, row, 32, item.resolution, note=Notes.PENDING_RESOLUTION)
    _set_cell(ws, row, 33, item.owner, note=Notes.SIGNOFF_ASSIGNED)
    _set_cell(ws, row, 34, item.priority)

