"""Extract structural specs from Deloitte WBS docx and RAID xlsx templates."""

from __future__ import annotations

from pathlib import Path
from typing import Any


def extract_wbs_spec(docx_path: Path) -> dict[str, Any]:
    from docx import Document

    document = Document(docx_path)
    if not document.tables:
        raise ValueError(f"No table found in WBS template: {docx_path}")

    table = document.tables[0]
    headers = [cell.text.strip() for cell in table.rows[0].cells]
    field_map = {
        "Level": ("level", "integer", "1=phase, 2=work stream, 3=task"),
        "WBS Code": ("wbs_code", "string", "Hierarchical code e.g. 1.0, 1.1, 1.1.1"),
        "WBS Element Title": ("wbs_element_title", "string", "Short name of the element"),
        "Description / Scope of Work": (
            "description_scope",
            "string",
            "Scope narrative for this element",
        ),
        "Key Deliverable": (
            "key_deliverable",
            "string",
            "Primary deliverable for this element",
        ),
    }
    columns: list[dict[str, str]] = []
    for header in headers:
        if header in field_map:
            field, typ, desc = field_map[header]
            columns.append({"header": header, "field": field, "type": typ, "description": desc})
    columns.append(
        {
            "header": "Owner Side",
            "field": "owner_side",
            "type": "string",
            "description": "Client, Consultant, or Joint responsibility for this row",
        }
    )

    example_rows: list[dict[str, Any]] = []
    for row in table.rows[1:8]:
        cells = [cell.text.strip() for cell in row.cells]
        if not any(cells):
            continue
        row_data = dict(zip(headers, cells, strict=False))
        example_rows.append(
            {
                "level": int(row_data.get("Level", "0") or 0),
                "wbs_code": row_data.get("WBS Code", ""),
                "wbs_element_title": row_data.get("WBS Element Title", ""),
                "description_scope": row_data.get("Description / Scope of Work", ""),
                "key_deliverable": row_data.get("Key Deliverable", ""),
                "owner_side": "Joint",
            }
        )

    return {
        "source_file": docx_path.name,
        "summary": (
            "Work packages must follow the Deloitte WBS table: a 3-level hierarchy "
            "(levels 1–3) with codes like 1.0, 1.1, 1.1.1. Each row is one element."
        ),
        "columns": columns,
        "hierarchy_rules": [
            "Level 1 = phase / major deliverable (WBS code X.0).",
            "Level 2 = work stream under that phase (X.Y).",
            "Level 3 = task under a work stream (X.Y.Z).",
            "parent_wbs_code must reference an existing wbs_code one level above.",
            "Minimum 3 level-1 elements, each with at least 2 level-2 children.",
            "Also include owner, estimated_duration, and dependencies (list of wbs_codes).",
            "Also include owner_side for each row: Client, Consultant, or Joint.",
        ],
        "example_rows": example_rows,
    }


def extract_raid_spec(xlsx_path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    workbook = load_workbook(xlsx_path, data_only=True)
    if "RAID Log - Detail" not in workbook.sheetnames:
        raise ValueError(f"RAID Log - Detail sheet not found in {xlsx_path}")
    detail = workbook["RAID Log - Detail"]

    sections = [
        {
            "key": "risks",
            "title": "Risks",
            "minimum": 3,
            "fields": [
                {"name": "id", "type": "string", "required": True, "description": "R1, R2, …"},
                {"name": "title", "type": "string", "required": True, "description": ""},
                {"name": "description", "type": "string", "required": True, "description": ""},
                {"name": "status", "type": "string", "required": True, "description": "e.g. Open"},
                {
                    "name": "mitigation",
                    "type": "string",
                    "required": True,
                    "description": "Mitigation plan",
                },
                {"name": "owner", "type": "string", "required": True, "description": ""},
                {
                    "name": "owner_side",
                    "type": "string",
                    "required": True,
                    "description": "Client, Consultant, or Joint",
                },
                {
                    "name": "wbs_links",
                    "type": "array",
                    "required": True,
                    "description": "WBS codes this risk relates to",
                },
                {
                    "name": "probability",
                    "type": "string",
                    "required": True,
                    "description": "High, Moderate, Low, or Very Low",
                },
                {
                    "name": "impact",
                    "type": "string",
                    "required": True,
                    "description": "High, Moderate, or Low",
                },
                {
                    "name": "risk_rating",
                    "type": "string",
                    "required": True,
                    "description": "Derived overall rating: High, Moderate, or Low",
                },
                {
                    "name": "risk_appetite",
                    "type": "string",
                    "required": True,
                    "description": "Engagement risk appetite (e.g. Moderate)",
                },
                {
                    "name": "priority",
                    "type": "string",
                    "required": True,
                    "description": "Response level: Monitor closely, Monitor, Stay informed, or Respond as required",
                },
            ],
            "template_headers": _row_headers(detail, 6, 1, 12),
        },
        {
            "key": "actions",
            "title": "Actions",
            "minimum": 2,
            "fields": [
                {"name": "id", "type": "string", "required": True, "description": "A1, A2, …"},
                {"name": "title", "type": "string", "required": True, "description": ""},
                {"name": "description", "type": "string", "required": True, "description": ""},
                {
                    "name": "status_update",
                    "type": "string",
                    "required": False,
                    "description": "MM-DD-YYYY: update text",
                },
                {"name": "status", "type": "string", "required": True, "description": ""},
                {"name": "owner", "type": "string", "required": True, "description": ""},
                {
                    "name": "owner_side",
                    "type": "string",
                    "required": True,
                    "description": "Client, Consultant, or Joint",
                },
                {
                    "name": "wbs_links",
                    "type": "array",
                    "required": True,
                    "description": "WBS codes",
                },
                {
                    "name": "priority",
                    "type": "string",
                    "required": True,
                    "description": "High, Moderate, or Low",
                },
            ],
            "template_headers": _row_headers(detail, 6, 14, 20),
        },
        {
            "key": "issues",
            "title": "Issues",
            "minimum": 0,
            "fields": [
                {"name": "id", "type": "string", "required": True, "description": "I1, I2, …"},
                {"name": "title", "type": "string", "required": True, "description": ""},
                {"name": "description", "type": "string", "required": True, "description": ""},
                {
                    "name": "resolution",
                    "type": "string",
                    "required": False,
                    "description": "Resolution when known",
                },
                {"name": "owner", "type": "string", "required": True, "description": ""},
                {
                    "name": "owner_side",
                    "type": "string",
                    "required": True,
                    "description": "Client, Consultant, or Joint",
                },
                {
                    "name": "priority",
                    "type": "string",
                    "required": True,
                    "description": "High, Moderate, or Low",
                },
            ],
            "template_headers": _row_headers(detail, 6, 22, 27),
        },
        {
            "key": "decisions",
            "title": "Decisions",
            "minimum": 0,
            "fields": [
                {"name": "id", "type": "string", "required": True, "description": "D1, D2, …"},
                {"name": "title", "type": "string", "required": True, "description": ""},
                {
                    "name": "description_justification",
                    "type": "string",
                    "required": True,
                    "description": "Description / justification",
                },
                {
                    "name": "resolution",
                    "type": "string",
                    "required": False,
                    "description": "Decision outcome",
                },
                {"name": "owner", "type": "string", "required": True, "description": ""},
                {
                    "name": "owner_side",
                    "type": "string",
                    "required": True,
                    "description": "Client, Consultant, or Joint",
                },
                {
                    "name": "priority",
                    "type": "string",
                    "required": True,
                    "description": "High, Moderate, or Low",
                },
            ],
            "template_headers": _row_headers(detail, 6, 29, 34),
        },
    ]

    return {
        "source_file": xlsx_path.name,
        "summary": (
            "RAID log follows the Deloitte Excel template on a single detail sheet: "
            "Risks, Actions, Issues, and Decisions are side-by-side sections. "
            "Risks include probability, impact, risk rating, risk appetite, and a "
            "response priority. Actions, issues, and decisions use High, Moderate, "
            "or Low priority. Each item labels whether the owner is on the Client, "
            "Consultant, or Joint side."
        ),
        "sections": sections,
    }


def _row_headers(ws: object, row: int, start_col: int, end_col: int) -> list[str]:
    from openpyxl.worksheet.worksheet import Worksheet

    assert isinstance(ws, Worksheet)
    headers: list[str] = []
    for col in range(start_col, end_col + 1):
        value = ws.cell(row=row, column=col).value
        if value is not None and str(value).strip():
            headers.append(str(value).strip())
    return headers


_RACI_SHEET = "RACI Table Sample"
_RACI_ACTIVITY_HEADER = "Deliverable / Activity"
_RACI_HEADER_ROW_LABEL = _RACI_ACTIVITY_HEADER
_RACI_LETTERS = frozenset({"R", "A", "C", "I"})


def extract_raci_spec(xlsx_path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    workbook = load_workbook(xlsx_path, data_only=True)
    if _RACI_SHEET not in workbook.sheetnames:
        raise ValueError(f"{_RACI_SHEET!r} sheet not found in {xlsx_path}")
    ws = workbook[_RACI_SHEET]

    header_row = _find_row_with_label(ws, _RACI_HEADER_ROW_LABEL, max_row=20)
    headers = _row_text_headers(ws, header_row)
    if not headers or headers[0] != _RACI_ACTIVITY_HEADER:
        raise ValueError(f"Expected {_RACI_ACTIVITY_HEADER!r} as first column in {xlsx_path}")
    roles = headers[1:]
    header_cols = _header_column_map(ws, header_row, headers)

    example_rows: list[dict[str, Any]] = []
    for row_idx in range(header_row + 1, ws.max_row + 1):
        row_data = _row_dict_by_columns(ws, row_idx, header_cols)
        activity = row_data.get(_RACI_ACTIVITY_HEADER, "").strip()
        if not activity:
            continue
        assignments: dict[str, str] = {}
        for role in roles:
            value = row_data.get(role, "").strip().upper()
            if value in _RACI_LETTERS:
                assignments[role] = value
        example_rows.append({"activity": activity, "assignments": assignments})

    return {
        "source_file": xlsx_path.name,
        "sheet": _RACI_SHEET,
        "summary": (
            "RACI matrix follows the Deloitte Excel template: rows are deliverables or "
            "activities from the WBS; columns are engagement roles; each cell is R, A, C, "
            "or I. Derive roles from the brief and WBS owners; align activities to level-2 "
            "or level-3 WBS elements."
        ),
        "activity_column": _RACI_ACTIVITY_HEADER,
        "roles": roles,
        "assignment_values": ["R", "A", "C", "I"],
        "generation_rules": [
            'Root JSON: {"roles": [ /* column headers */ ], "rows": [ /* one object per activity row */ ]}.',
            "Each row object: activity (string) and assignments (object mapping role name to R, A, C, or I).",
            "Every role in roles must appear as a key in each row's assignments.",
            "Each activity row should have exactly one Accountable (A) assignment.",
            "Use R for Responsible, A for Accountable, C for Consulted, I for Informed.",
            "Ground roles in named stakeholders and WBS owners from the brief; do not invent generic columns.",
            "Include at least one row per major WBS deliverable or work stream.",
        ],
        "example_rows": example_rows,
    }


_STATUS_SECTION_MAP: list[tuple[str, str, str]] = [
    ("Summary:", "summary", "High-level status; reference RAID decisions and progress"),
    ("New Initiatives:", "new_initiatives", "New actions from RAID; list of strings"),
    ("Blockers:", "blockers", "Items needing management help; from RAID issues or WBS dependencies"),
    ("Risks:", "risks", "High-probability/high-impact risks from RAID; list of strings"),
    ("Milestones:", "milestones", "Milestone table aligned to WBS timeline"),
    ("Details:", "supporting_links", "Links to supporting documentation; list of strings"),
]


def extract_status_report_spec(docx_path: Path) -> dict[str, Any]:
    from docx import Document

    document = Document(docx_path)
    sections: list[dict[str, str]] = []
    for label, field, description in _STATUS_SECTION_MAP:
        sections.append({"label": label, "field": field, "description": description})

    header_fields = [
        {
            "field": "project_name",
            "label": "Project Name:",
            "type": "string",
            "description": "Engagement or project title from the brief",
        },
        {
            "field": "project_manager",
            "label": "Project Manager:",
            "type": "string",
            "description": "Named PM or lead from the brief when stated",
        },
        {
            "field": "date",
            "label": "Date:",
            "type": "string",
            "description": "Week ending or report date (ISO or brief format)",
        },
    ]

    milestone_columns: list[dict[str, str]] = []
    example_milestones: list[dict[str, str]] = []
    if document.tables:
        table = document.tables[0]
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        field_map = {
            "Milestone": ("name", "Milestone name aligned to WBS"),
            "Target date": ("target_date", "Planned completion date"),
            "Actual date": ("actual_date", "Actual completion date when known"),
        }
        for header in headers:
            if header in field_map:
                field, desc = field_map[header]
                milestone_columns.append({"header": header, "field": field, "description": desc})
        milestone_columns.append(
            {
                "header": "Owner Side",
                "field": "owner_side",
                "description": "Client, Consultant, or Joint accountable side",
            }
        )
        for row in table.rows[1:4]:
            cells = [cell.text.strip() for cell in row.cells]
            if not any(cells):
                continue
            row_data = dict(zip(headers, cells, strict=False))
            example_milestones.append(
                {
                    "name": row_data.get("Milestone", ""),
                    "target_date": row_data.get("Target date", ""),
                    "actual_date": row_data.get("Actual date") or None,
                    "owner_side": "Joint",
                    "status": "Not Started",
                }
            )

    return {
        "source_file": docx_path.name,
        "summary": (
            "Weekly status report follows the Deloitte Word template: header fields, "
            "narrative summary, bullet sections (initiatives, blockers, risks), milestone "
            "table, and supporting links. Derive content only from brief, WBS, and RAID."
        ),
        "header_fields": header_fields,
        "sections": sections,
        "milestone_columns": milestone_columns,
        "milestone_status_values": [
            "Not Started",
            "In Progress",
            "Completed",
            "Delayed",
        ],
        "generation_rules": [
            "Root JSON includes project_name, project_manager, date, summary, "
            "new_initiatives, blockers, risks, milestones, supporting_links.",
            "new_initiatives, blockers, risks, and supporting_links are arrays of strings.",
            "Each milestone object: name, target_date, actual_date (optional), owner_side, status.",
            "owner_side must be Client, Consultant, or Joint.",
            "Risks and blockers must trace to RAID; milestones must align with WBS.",
        ],
        "example_milestones": example_milestones,
    }


def _find_row_with_label(ws: object, label: str, *, max_row: int) -> int:
    from openpyxl.worksheet.worksheet import Worksheet

    assert isinstance(ws, Worksheet)
    for row in range(1, max_row + 1):
        for col in range(1, ws.max_column + 1):
            value = ws.cell(row=row, column=col).value
            if value is not None and str(value).strip() == label:
                return row
    raise ValueError(f"Row with label {label!r} not found")


def _row_text_headers(ws: object, row: int) -> list[str]:
    from openpyxl.worksheet.worksheet import Worksheet

    assert isinstance(ws, Worksheet)
    headers: list[str] = []
    for col in range(1, ws.max_column + 1):
        value = ws.cell(row=row, column=col).value
        if value is not None and str(value).strip():
            headers.append(str(value).strip())
    return headers


def _header_column_map(ws: object, row: int, headers: list[str]) -> dict[str, int]:
    from openpyxl.worksheet.worksheet import Worksheet

    assert isinstance(ws, Worksheet)
    mapping: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        value = ws.cell(row=row, column=col).value
        if value is not None:
            text = str(value).strip()
            if text in headers:
                mapping[text] = col
    return mapping


def _row_dict_by_columns(
    ws: object, row: int, header_cols: dict[str, int]
) -> dict[str, str]:
    from openpyxl.worksheet.worksheet import Worksheet

    assert isinstance(ws, Worksheet)
    row_data: dict[str, str] = {}
    for header, col in header_cols.items():
        value = ws.cell(row=row, column=col).value
        row_data[header] = str(value).strip() if value is not None else ""
    return row_data
