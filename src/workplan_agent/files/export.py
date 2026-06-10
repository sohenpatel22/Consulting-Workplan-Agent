"""Export markdown artifacts to downloadable formats."""

from __future__ import annotations

import io
import re
from html import escape
from typing import Literal

from workplan_agent.models.brief import ParsedBrief
from workplan_agent.models.raid import RAIDLog
from workplan_agent.models.raci import RaciMatrix
from workplan_agent.models.status_report import WeeklyStatusReport
from workplan_agent.models.wbs import WBS

ExportFormat = Literal["md", "html", "docx", "pdf", "xlsx"]
ArtifactKind = Literal["wbs", "raid", "raci", "status"]

CONTENT_TYPES: dict[ExportFormat, str] = {
    "md": "text/markdown; charset=utf-8",
    "html": "text/html; charset=utf-8",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


class ExportError(ValueError):
    """Raised when export cannot be produced."""


def export_artifact(
    markdown: str,
    artifact: ArtifactKind,
    fmt: ExportFormat,
    *,
    wbs: WBS | None = None,
    raid: RAIDLog | None = None,
    raci_matrix: RaciMatrix | None = None,
    status_report: WeeklyStatusReport | None = None,
    parsed_brief: ParsedBrief | None = None,
    use_deloitte_templates: bool = True,
) -> tuple[bytes, str]:
    """Return (file_bytes, suggested_filename) for the given markdown and format."""
    base_name = artifact

    if fmt == "xlsx":
        if parsed_brief is None:
            raise ExportError("Parsed brief is required for XLSX export.")
        if artifact == "raid":
            if raid is None:
                raise ExportError("RAID structured data is required for XLSX export.")
            from workplan_agent.files.deloitte_export import export_raid_deloitte_xlsx

            return export_raid_deloitte_xlsx(raid, parsed_brief), f"{base_name}.xlsx"
        if artifact == "raci":
            if raci_matrix is None:
                raise ExportError("RACI structured data is required for XLSX export.")
            from workplan_agent.files.deloitte_export import export_raci_deloitte_xlsx

            return export_raci_deloitte_xlsx(raci_matrix, parsed_brief), f"{base_name}.xlsx"
        raise ExportError("XLSX export is only available for RAID and RACI.")

    if fmt == "docx" and artifact == "wbs" and use_deloitte_templates and wbs and parsed_brief:
        from workplan_agent.files.deloitte_export import export_wbs_deloitte_docx

        return export_wbs_deloitte_docx(wbs, parsed_brief), f"{base_name}.docx"

    if fmt == "pdf" and use_deloitte_templates and parsed_brief:
        from workplan_agent.files import deloitte_export as de

        if artifact == "wbs" and wbs is not None:
            return de.export_wbs_deloitte_pdf(wbs, parsed_brief), f"{base_name}.pdf"
        if artifact == "raid" and raid is not None:
            return de.export_raid_deloitte_pdf(raid, parsed_brief), f"{base_name}.pdf"
        if artifact == "raci" and raci_matrix is not None:
            return de.export_raci_deloitte_pdf(
                raci_matrix, parsed_brief
            ), f"{base_name}.pdf"
        if artifact == "status" and status_report is not None:
            return de.export_status_deloitte_pdf(status_report, parsed_brief), f"{base_name}.pdf"

    if not markdown.strip():
        raise ExportError(f"No {artifact.upper()} content available to export.")

    if fmt == "md":
        return markdown.encode("utf-8"), f"{base_name}.md"
    if fmt == "html":
        return _to_html(markdown).encode("utf-8"), f"{base_name}.html"
    if fmt == "docx":
        return _to_docx(markdown), f"{base_name}.docx"
    if fmt == "pdf":
        return _to_pdf(markdown), f"{base_name}.pdf"
    raise ExportError(f"Unsupported export format: {fmt}")


def markdown_to_html_body(markdown: str) -> str:
    """Convert markdown to an HTML fragment (for in-app preview)."""
    try:
        import markdown as md_lib
    except ImportError as exc:
        raise ExportError("HTML export is not installed.") from exc

    return md_lib.markdown(
        markdown,
        extensions=["tables", "fenced_code", "nl2br"],
    )


def _to_html(markdown: str) -> str:
    body = markdown_to_html_body(markdown)
    title_match = re.search(r"^#\s+(.+)$", markdown, re.MULTILINE)
    title = escape(title_match.group(1).strip()) if title_match else "Workplan artifact"
    return f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>{title}</title>
  <style>
    body {{ font-family: system-ui, sans-serif; line-height: 1.5; max-width: 52rem; margin: 2rem auto; padding: 0 1rem; color: #1a1a1a; }}
    table {{ border-collapse: collapse; width: 100%; margin: 1rem 0; }}
    th, td {{ border: 1px solid #ccc; padding: 0.4rem 0.6rem; text-align: left; }}
    th {{ background: #f4f4f4; }}
    pre {{ background: #f6f8fa; padding: 0.75rem; overflow-x: auto; border-radius: 6px; }}
    code {{ font-family: ui-monospace, monospace; }}
  </style>
</head>
<body>
{body}
</body>
</html>
"""


def _parse_table_cells(line: str) -> list[str] | None:
    if not (line.startswith("|") and line.endswith("|")):
        return None
    cells = [c.strip() for c in line.strip("|").split("|")]
    if _is_table_separator_cells(cells):
        return None
    return cells


def _is_table_separator_line(line: str) -> bool:
    stripped = line.strip()
    if not (stripped.startswith("|") and stripped.endswith("|")):
        return False
    cells = [cell.strip() for cell in stripped.strip("|").split("|")]
    return _is_table_separator_cells(cells)


def _is_table_separator_cells(cells: list[str]) -> bool:
    return bool(cells) and all(set(cell) <= {"-", ":"} for cell in cells)


def _to_docx(markdown: str) -> bytes:
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError as exc:
        raise ExportError("DOCX export is not installed.") from exc

    document = Document()
    pending_table_rows: list[list[str]] = []

    def flush_table() -> None:
        if not pending_table_rows:
            return
        table = document.add_table(rows=len(pending_table_rows), cols=len(pending_table_rows[0]))
        table.style = "Table Grid"
        for row_idx, row_cells in enumerate(pending_table_rows):
            for col_idx, cell_text in enumerate(row_cells):
                table.rows[row_idx].cells[col_idx].text = cell_text
        pending_table_rows.clear()

    for line in markdown.splitlines():
        stripped = line.rstrip()
        if not stripped:
            flush_table()
            continue

        table_cells = _parse_table_cells(stripped)
        if table_cells is not None:
            pending_table_rows.append(table_cells)
            continue

        flush_table()
        if stripped.startswith("# "):
            document.add_heading(stripped[2:].strip(), level=1)
        elif stripped.startswith("## "):
            document.add_heading(stripped[3:].strip(), level=2)
        elif stripped.startswith("### "):
            document.add_heading(stripped[4:].strip(), level=3)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            document.add_paragraph(stripped[2:].strip(), style="List Bullet")
        else:
            para = document.add_paragraph(stripped)
            for run in para.runs:
                run.font.size = Pt(11)

    flush_table()
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _strip_inline_markdown(text: str) -> str:
    """Remove simple inline markdown markers for PDF output."""
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"__(.+?)__", r"\1", text)
    text = re.sub(r"_(.+?)_", r"\1", text)
    return text.strip()


def _pdf_safe(text: str) -> str:
    """Helvetica core font is Latin-1; normalize unsupported characters."""
    if not text:
        return ""
    normalized = (
        text.replace("\u2014", " - ")
        .replace("\u2013", "-")
        .replace("\u2022", "*")
        .replace("\u2018", "'")
        .replace("\u2019", "'")
        .replace("\u201c", '"')
        .replace("\u201d", '"')
    )
    return normalized.encode("latin-1", errors="replace").decode("latin-1")


def _pdf_plain_line(line: str) -> str:
    """Convert a markdown line to plain text for paragraph/heading rendering."""
    stripped = line.rstrip()
    if stripped.startswith("# "):
        return _strip_inline_markdown(stripped[2:].strip())
    if stripped.startswith("## "):
        return _strip_inline_markdown(stripped[3:].strip())
    if stripped.startswith("### "):
        return _strip_inline_markdown(stripped[4:].strip())
    if stripped.startswith("- ") or stripped.startswith("* "):
        return f"- {_strip_inline_markdown(stripped[2:].strip())}"
    return _strip_inline_markdown(stripped)


def _equal_col_widths(column_count: int, total_width: float) -> tuple[float, ...]:
    width = total_width / max(column_count, 1)
    return tuple(width for _ in range(column_count))


def _write_pdf_table(
    pdf: object,
    rows: list[list[str]],
    *,
    landscape: bool = False,
    new_page: bool = True,
) -> None:
    from fpdf.fonts import FontFace

    if not rows:
        return

    if new_page:
        if landscape:
            pdf.add_page(orientation="L")  # type: ignore[attr-defined]
        else:
            pdf.add_page()  # type: ignore[attr-defined]

    width = pdf.epw  # type: ignore[attr-defined]
    col_count = len(rows[0])
    col_widths = _equal_col_widths(col_count, width)
    font_size = 7 if col_count >= 6 else 9
    line_height = 4 if col_count >= 6 else 5
    header_style = FontFace(emphasis="BOLD")

    pdf.set_font("Helvetica", size=font_size)  # type: ignore[attr-defined]
    with pdf.table(  # type: ignore[attr-defined]
        width=width,
        col_widths=col_widths,
        line_height=line_height,
        text_align="LEFT",
    ) as table:
        for row_idx, cells in enumerate(rows):
            row = table.row()
            for cell in cells:
                text = _pdf_safe(_strip_inline_markdown(cell))
                if row_idx == 0:
                    row.cell(text, style=header_style)
                else:
                    row.cell(text)


def _to_pdf(markdown: str) -> bytes:
    try:
        from fpdf import FPDF
    except ImportError as exc:
        raise ExportError("PDF export is not installed.") from exc

    pdf = FPDF()
    pdf.set_margins(18, 18, 18)
    pdf.set_auto_page_break(auto=True, margin=18)
    width = pdf.epw
    pending_table_rows: list[list[str]] = []
    content_started = False

    def ensure_page() -> None:
        nonlocal content_started
        if not content_started:
            pdf.add_page()
            content_started = True

    def flush_table(*, continue_on_portrait: bool = True) -> None:
        nonlocal pending_table_rows, content_started
        if not pending_table_rows:
            return
        use_landscape = len(pending_table_rows[0]) >= 6
        if not content_started:
            if use_landscape:
                pdf.add_page(orientation="L")
            else:
                pdf.add_page()
            content_started = True
            _write_pdf_table(
                pdf,
                pending_table_rows,
                landscape=use_landscape,
                new_page=False,
            )
        else:
            _write_pdf_table(pdf, pending_table_rows, landscape=use_landscape)
        pending_table_rows = []
        if continue_on_portrait:
            pdf.add_page()

    for line in markdown.splitlines():
        stripped = line.rstrip()
        if not stripped:
            flush_table()
            if content_started:
                pdf.ln(4)
            continue

        if _is_table_separator_line(stripped):
            continue

        table_cells = _parse_table_cells(stripped)
        if table_cells is not None:
            pending_table_rows.append(table_cells)
            continue

        flush_table()
        ensure_page()

        if stripped.startswith("# "):
            pdf.set_font("Helvetica", style="B", size=16)
            pdf.multi_cell(width, 8, _pdf_safe(_pdf_plain_line(stripped)))
            pdf.set_font("Helvetica", style="", size=11)
            continue
        if stripped.startswith("## "):
            pdf.set_font("Helvetica", style="B", size=13)
            pdf.multi_cell(width, 7, _pdf_safe(_pdf_plain_line(stripped)))
            pdf.set_font("Helvetica", style="", size=11)
            continue
        if stripped.startswith("### "):
            pdf.set_font("Helvetica", style="B", size=12)
            pdf.multi_cell(width, 6, _pdf_safe(_pdf_plain_line(stripped)))
            pdf.set_font("Helvetica", style="", size=11)
            continue

        text = _pdf_safe(_pdf_plain_line(stripped))
        if not text:
            continue
        pdf.multi_cell(width, 5, text)

    flush_table(continue_on_portrait=False)
    return bytes(pdf.output())
