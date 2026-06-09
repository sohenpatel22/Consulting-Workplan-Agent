import io

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter

from workplan_agent.api.app import app
from workplan_agent.api.store import reset_store
from workplan_agent.files.deloitte_export import export_raid_deloitte_xlsx, export_wbs_deloitte_docx
from workplan_agent.files.export import ExportError, export_artifact
from workplan_agent.render import (
    render_raid_markdown,
    render_raci_markdown,
    render_status_markdown,
    render_wbs_markdown,
)
from workplan_agent.files.extract import BriefExtractError, extract_text_from_bytes
from workplan_agent.llm.fixtures import (
    PARSED_BRIEF_FIXTURE,
    RAID_FIXTURE,
    RACI_MATRIX_FIXTURE,
    STATUS_REPORT_FIXTURE,
    WBS_FIXTURE,
)
from workplan_agent.models.brief import ParsedBrief
from workplan_agent.models.raid import RAIDLog
from workplan_agent.models.raci import RaciMatrix
from workplan_agent.models.status_report import WeeklyStatusReport
from workplan_agent.models.wbs import WBS

BRIEF = """# Test Brief

**Client:** Acme Corp
**Project:** Platform Rollout

## Objectives

1. Deliver MVP
2. Train operations team
"""

SAMPLE_MD = """# Sample

## Section

Some body text for extraction testing purposes.
"""


@pytest.fixture
def client() -> TestClient:
    reset_store()
    return TestClient(app)


def test_extract_markdown_bytes() -> None:
    text = extract_text_from_bytes(SAMPLE_MD.encode(), "brief.md")
    assert "Sample" in text
    assert len(text) >= 10


def test_extract_rejects_unknown_extension() -> None:
    with pytest.raises(BriefExtractError, match="Unsupported"):
        extract_text_from_bytes(b"data", "file.exe")


def test_extract_docx_bytes() -> None:
    buffer = io.BytesIO()
    doc = Document()
    doc.add_paragraph("Client brief paragraph for upload testing here.")
    doc.save(buffer)
    text = extract_text_from_bytes(buffer.getvalue(), "brief.docx")
    assert "Client brief paragraph" in text


def test_extract_pdf_bytes() -> None:
    buffer = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.write(buffer)
    with pytest.raises(BriefExtractError, match="too short"):
        extract_text_from_bytes(buffer.getvalue(), "brief.pdf")


def test_export_formats() -> None:
    md = "# WBS\n\n## Phase 1\n\n| ID | Task |\n| --- | --- |\n| 1.1 | Setup |\n"
    for fmt in ("md", "html", "docx", "pdf"):
        content, name = export_artifact(md, "wbs", fmt)
        assert content
        assert name.endswith(f".{fmt}")


def test_deloitte_wbs_docx_export() -> None:
    wbs = WBS.model_validate(WBS_FIXTURE)
    brief = ParsedBrief.model_validate(PARSED_BRIEF_FIXTURE)
    content = export_wbs_deloitte_docx(wbs, brief)
    assert content[:2] == b"PK"


def test_deloitte_wbs_pdf_export() -> None:
    from pypdf import PdfReader

    wbs = WBS.model_validate(WBS_FIXTURE)
    brief = ParsedBrief.model_validate(PARSED_BRIEF_FIXTURE)
    md = render_wbs_markdown(wbs, brief)
    content, name = export_artifact(md, "wbs", "pdf", wbs=wbs, parsed_brief=brief)
    assert name == "wbs.pdf"
    assert content.startswith(b"%PDF")
    assert len(content) > 2_500
    text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages)
    assert "Discovery" in text
    assert "Stakeholder interviews" in text
    assert "1.0" in text


def test_deloitte_raid_pdf_export() -> None:
    from pypdf import PdfReader

    raid = RAIDLog.model_validate(RAID_FIXTURE)
    brief = ParsedBrief.model_validate(PARSED_BRIEF_FIXTURE)
    md = render_raid_markdown(raid, brief)
    content, name = export_artifact(md, "raid", "pdf", raid=raid, parsed_brief=brief)
    assert name == "raid.pdf"
    text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages)
    assert "SME availability" in text
    assert "|----|" not in text
    assert "Risks" in text


def test_deloitte_raci_pdf_export() -> None:
    from pypdf import PdfReader

    raci = RaciMatrix.model_validate(RACI_MATRIX_FIXTURE)
    brief = ParsedBrief.model_validate(PARSED_BRIEF_FIXTURE)
    md = render_raci_markdown(raci, brief)
    content, name = export_artifact(
        md, "raci", "pdf", raci_matrix=raci, parsed_brief=brief
    )
    assert name == "raci.pdf"
    text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages)
    assert "Stakeholder interviews" in text
    assert "1.1" in text
    assert "|-------------|" not in text


def test_deloitte_status_pdf_export() -> None:
    from pypdf import PdfReader

    status = WeeklyStatusReport.model_validate(STATUS_REPORT_FIXTURE)
    brief = ParsedBrief.model_validate(PARSED_BRIEF_FIXTURE)
    md = render_status_markdown(status, brief)
    content, name = export_artifact(
        md, "status", "pdf", status_report=status, parsed_brief=brief
    )
    assert name == "status.pdf"
    text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages)
    assert "Weekly Status Report" in text
    assert "kickoff" in text.lower()
    assert "Architecture approved" in text


def test_deloitte_raid_xlsx_export() -> None:
    from io import BytesIO

    from openpyxl import load_workbook

    raid = RAIDLog.model_validate(RAID_FIXTURE)
    brief = ParsedBrief.model_validate(PARSED_BRIEF_FIXTURE)
    content = export_raid_deloitte_xlsx(raid, brief)
    assert content[:2] == b"PK"
    workbook = load_workbook(BytesIO(content))
    assert workbook.sheetnames == ["RAID Log - Detail"]
    ws = workbook["RAID Log - Detail"]
    assert ws.cell(row=7, column=1).value == "R1"
    assert ws.cell(row=7, column=9).value == "Moderate"
    assert ws.cell(row=7, column=10).value == "High"
    assert ws.cell(row=7, column=12).value == "Monitor closely"
    assert ws.cell(row=7, column=14).value == "A1"
    assert ws.cell(row=7, column=22).value is None
    assert ws.cell(row=7, column=29).value == "D1"


def test_deloitte_raid_xlsx_export_many_rows() -> None:
    """Detail sheet has merged cells in lower rows; export must not raise."""
    import copy

    data = copy.deepcopy(RAID_FIXTURE)
    for index in range(4, 16):
        data["risks"].append(
            {**data["risks"][0], "id": f"R{index}", "title": f"Extra risk {index}"}
        )
    raid = RAIDLog.model_validate(data)
    brief = ParsedBrief.model_validate(PARSED_BRIEF_FIXTURE)
    content = export_raid_deloitte_xlsx(raid, brief)
    assert content[:2] == b"PK"


def test_export_raid_xlsx_requires_models() -> None:
    with pytest.raises(ExportError, match="Parsed brief"):
        export_artifact("# RAID", "raid", "xlsx")


def test_upload_endpoint_md(client: TestClient) -> None:
    response = client.post(
        "/api/brief/upload",
        files={"file": ("brief.md", SAMPLE_MD.encode(), "text/markdown")},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["filename"] == "brief.md"
    assert "Sample" in data["text"]


def test_download_requires_session(client: TestClient) -> None:
    response = client.get("/api/sessions/missing/download/wbs/md")
    assert response.status_code == 404


def test_download_after_generation(client: TestClient) -> None:
    import time

    create = client.post("/api/sessions", json={"brief_text": BRIEF, "mock": True})
    session_id = create.json()["session_id"]

    for _ in range(300):
        session = client.get(f"/api/sessions/{session_id}").json()
        if session.get("awaiting_confirmation") or session.get("phase") == "awaiting_confirmation":
            # Confirm open questions
            questions = session.get("confirmation_questions") or []
            answers = [
                {
                    "question": q if isinstance(q, str) else q.get("question", ""),
                    "answer": f"Confirmed for planning: {q if isinstance(q, str) else q.get('question', '')}",
                }
                for q in questions
            ]
            client.post(
                f"/api/sessions/{session_id}/confirm",
                json={"answers": answers},
            )
            continue
        if session["job_status"] in ("complete", "failed"):
            break
        time.sleep(0.05)

    assert session["job_status"] == "complete"

    for fmt, content_prefix in (
        ("md", b"#"),
        ("html", b"<!DOCTYPE html>"),
        ("docx", b"PK"),
        ("pdf", b"%PDF"),
    ):
        response = client.get(f"/api/sessions/{session_id}/download/wbs/{fmt}")
        assert response.status_code == 200, (fmt, response.text)
        assert response.content.startswith(content_prefix)

    xlsx = client.get(f"/api/sessions/{session_id}/download/raid/xlsx")
    assert xlsx.status_code == 200
    assert xlsx.content.startswith(b"PK")
